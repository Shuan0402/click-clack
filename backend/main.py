import os
import time
import io
import math
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from openai import OpenAI
import pypdf
from pptx import Presentation

load_dotenv()

api_key = os.getenv("OPENAI_API_KEY")
client = None

if api_key:
    client = OpenAI(
        api_key=api_key,
        base_url="https://api.groq.com/openai/v1"
    )
else:
    print("⚠️ 警告: 未設定 API Key")

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], 
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

class GenerateRequest(BaseModel):
    prompt: str
    mode: str = "creative" 
    length: str = "medium"
    language: str = "en"

async def extract_text_from_file(file: UploadFile) -> str:
    content = ""
    filename = file.filename.lower()
    file_bytes = await file.read()
    file_stream = io.BytesIO(file_bytes)
    try:
        if filename.endswith(".txt"):
            content = file_bytes.decode("utf-8")
        elif filename.endswith(".pdf"):
            pdf_reader = pypdf.PdfReader(file_stream)
            for page in pdf_reader.pages:
                text = page.extract_text()
                if text: content += text + "\n"
        elif filename.endswith(".pptx"):
            prs = Presentation(file_stream)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        content += shape.text + "\n"
    except Exception as e:
        print(f"解析失敗: {e}")
        return ""
    return content.strip()

def chunk_text(text, chunk_size=4000):
    return [text[i:i+chunk_size] for i in range(0, len(text), chunk_size)]

def get_language_instruction(lang: str):
    if lang == "zh-TW":
        return "Please respond ONLY in Traditional Chinese (Taiwan 繁體中文). Ensure the tone is natural for Taiwanese users."
    return "Please respond ONLY in English."

def get_target_length_desc(length_code, lang):
    word_count_map = {
        "short": {"en": "around 100 words", "zh": "約 150 個中文字"},
        "medium": {"en": "around 250 words", "zh": "約 350 個中文字"},
        "long": {"en": "around 500 words", "zh": "約 600 個中文字"}
    }
    config = word_count_map.get(length_code, word_count_map["medium"])
    return config["zh" if lang == "zh-TW" else "en"]

# --- API 1: 純文字生成 ---
@app.post("/api/generate")
async def generate_text(request: GenerateRequest):
    lang = request.language
    target_len = get_target_length_desc(request.length, lang) 
    lang_inst = get_language_instruction(lang)

    if client:
        try:
            sys_msg = f"You are a creative writer. Generate an article ({target_len}) for typing practice. {lang_inst}"
            completion = client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                messages=[{"role": "system", "content": sys_msg}, {"role": "user", "content": request.prompt}],
                temperature=0.7
            )
            return {"status": "success", "data": completion.choices[0].message.content}
        except Exception as e:
            print(f"API Error: {e}")
    return {"status": "success", "data": f"[Mock] Generated based on: {request.prompt}"}

# --- API 2: 檔案分析 ---
@app.post("/api/analyze")
async def analyze_file(
    file: UploadFile = File(...), 
    mode: str = Form(...),    
    length: str = Form(...),
    language: str = Form("en")
):
    raw_text = await extract_text_from_file(file)
    if not raw_text:
        raise HTTPException(status_code=400, detail="無法讀取檔案內容")
    
    total_chars = len(raw_text)
    target_len_desc = get_target_length_desc(length, language)
    lang_inst = get_language_instruction(language)

    if client:
        try:
            final_context = raw_text
            if total_chars > 4000:
                chunks = chunk_text(raw_text, 4000)
                summaries = []
                for chunk in chunks:
                    completion = client.chat.completions.create(
                        model="llama-3.3-70b-versatile",
                        messages=[
                            {"role": "system", "content": "Concise summarizer. Facts only."},
                            {"role": "user", "content": chunk}
                        ],
                        temperature=0.5,
                    )
                    summaries.append(completion.choices[0].message.content)
                final_context = "\n".join(summaries)

            if mode == "extract":
                final_system_prompt = f"Summarizer. Reconstruct article ({target_len_desc}). {lang_inst} Plain text only."
            else:
                final_system_prompt = f"Teacher. Expand with knowledge ({target_len_desc}). {lang_inst} For typing practice."

            final_response = client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                messages=[
                    {"role": "system", "content": final_system_prompt},
                    {"role": "user", "content": f"Context:\n\n{final_context}"}
                ],
                temperature=0.7,
            )
            return {"status": "success", "data": final_response.choices[0].message.content}
        except Exception as e:
            return {"status": "error", "message": str(e)}

    return {"status": "success", "data": "Mock data"}

@app.get("/")
def read_root():
    return {"message": "ClickClack Backend is ready!"}